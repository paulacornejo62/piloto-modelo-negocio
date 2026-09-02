#!/usr/bin/env python3
"""Genera la planilla de Proyección + Hitos comerciales desde el expediente.

Produce un .xlsx con tres hojas — Unit economics, Proyección (por escenario) e
Hitos comerciales — pre-pobladas desde `modelo_negocio` y `plan_escalamiento`.
Es un **insumo de arranque** que el consultor traspasa/valida al formato del
fondo. Requiere openpyxl (en el .venv del setup).

Uso: python3 render/gantt.py <tenant>   (mejor con .venv/bin/python)
"""
from __future__ import annotations

import json
import sys
from pathlib import Path

import openpyxl
from openpyxl.styles import Font

for _s in (sys.stdout, sys.stderr):  # consola Windows (cp1252) no soporta emojis/guiones
    try:
        _s.reconfigure(encoding="utf-8")
    except Exception:
        pass

SISTEMA = Path(__file__).resolve().parent.parent
NEG = Font(bold=True)


def construir(datos: dict, destino: Path) -> None:
    mn = datos.get("modelo_negocio", {}) or {}
    pe = datos.get("plan_escalamiento", {}) or {}
    ue = mn.get("unit_economics", {}) or {}
    wb = openpyxl.Workbook()

    # --- Hoja Unit economics ---
    ws = wb.active
    ws.title = "Unit economics"
    ws.append(["Indicador", "Valor", "Tipo"])
    for c in ws[1]:
        c.font = NEG
    ws.append(["Unidad del modelo", ue.get("unidad_modelo", ""), "—"])
    filas_ue = [
        ("Precio unitario (CLP)", "precio_unitario", "entrada"),
        ("Costo variable unitario (CLP)", "costo_variable_unitario", "entrada"),
        ("Costos fijos mensuales (CLP)", "costos_fijos_mensuales", "entrada"),
        ("CAC (CLP)", "cac_clp", "entrada"),
        ("Vida del cliente (meses)", "vida_cliente_meses", "entrada"),
        ("Margen de contribución unitario (CLP)", "margen_contribucion_unitario", "motor"),
        ("Margen de contribución %", "margen_contribucion_pct", "motor"),
        ("LTV (CLP)", "ltv_clp", "motor"),
        ("LTV / CAC", "ltv_cac_ratio", "motor"),
        ("Payback de CAC (meses)", "payback_cac_meses", "motor"),
        ("Punto de equilibrio (unidades)", "punto_equilibrio_unidades", "motor"),
        ("Punto de equilibrio (CLP)", "punto_equilibrio_clp", "motor"),
    ]
    for label, k, tipo in filas_ue:
        ws.append([label, ue.get(k), tipo])

    # --- Hoja Proyección ---
    pr = wb.create_sheet("Proyección")
    pr.append(["Escenario", "Año", "Unidades", "Ingresos (CLP)", "Costos (CLP)", "Resultado (CLP)"])
    for c in pr[1]:
        c.font = NEG
    escenarios = (mn.get("proyeccion", {}) or {}).get("escenarios", []) or []
    if escenarios:
        for e in escenarios:
            for a in e.get("anios", []) or []:
                pr.append([e.get("nombre", ""), a.get("anio", ""), a.get("unidades", ""),
                           a.get("ingresos_clp", ""), a.get("costos_clp", ""), a.get("resultado_clp", "")])
    else:
        pr.append(["(sin proyección en el expediente — completar modelo_negocio.proyeccion)"])

    # --- Hoja Hitos comerciales ---
    hc = wb.create_sheet("Hitos comerciales")
    hc.append(["N°", "Hito", "Meta (indicador medible)", "Fecha objetivo"])
    for c in hc[1]:
        c.font = NEG
    hitos = pe.get("hitos_comerciales", []) or []
    if hitos:
        for h in hitos:
            hc.append([h.get("n", ""), h.get("nombre", ""), h.get("meta", ""), h.get("fecha_objetivo", "")])
    else:
        hc.append(["(sin hitos en el expediente — completar plan_escalamiento.hitos_comerciales)"])
    if pe.get("metas_traccion"):
        hc.append([])
        hc.append(["Metas de tracción:", pe["metas_traccion"]])

    for sh in wb.worksheets:
        for col in sh.columns:
            sh.column_dimensions[col[0].column_letter].width = 24

    wb.save(destino)


def main() -> None:
    if len(sys.argv) < 2:
        sys.exit("Uso: python3 render/gantt.py <tenant>")
    tenant = sys.argv[1]
    ruta = SISTEMA / "proyectos" / tenant / "expediente.json"
    if not ruta.exists():
        sys.exit(f"ERROR: no existe {ruta}")
    destino = ruta.parent / "plan-trabajo.xlsx"
    construir(json.loads(ruta.read_text(encoding="utf-8")), destino)
    print(f"Generado: {destino} (insumo de arranque; traspasar al formato del fondo)")


if __name__ == "__main__":
    main()
